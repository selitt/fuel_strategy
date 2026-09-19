import csv
from io import BytesIO, StringIO

import openpyxl
import pytest
import xlwt


HEADERS = [
    "record_type", "plan_id", "scenario_id", "year", "source_id",
    "order_t", "reserved_capacity_t", "investment_id", "initial_inventory_t",
]


def rows_for(plan):
    rows = [
        ["plan", plan["plan_id"], plan["scenario_id"], "", "", "", "", "", ""],
        ["inventory_policy", "", "", "", "", "", "", "", plan["decisions"]["inventory_policy"]["initial_inventory_t"]],
    ]
    rows.extend([
        "supply_order", "", "", row["year"], row["source_id"], row["order_t"], "", "", "",
    ] for row in plan["decisions"]["supply_orders"])
    rows.extend([
        "capacity_reservation", "", "", row["year"], row["source_id"], "", row["reserved_capacity_t"], "", "",
    ] for row in plan["decisions"]["capacity_reservations"])
    rows.extend([
        "investment", "", "", row["year"], "", "", "", row["investment_id"], "",
    ] for row in plan["decisions"]["investments"])
    return rows


def csv_file(plan, *, delimiter=","):
    output = StringIO()
    writer = csv.writer(output, delimiter=delimiter)
    writer.writerow(HEADERS)
    writer.writerows(rows_for(plan))
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def xlsx_file(plan, *, formula=False):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "plan"
    sheet.append(HEADERS)
    for row in rows_for(plan):
        sheet.append(row)
    if formula:
        sheet.cell(row=4, column=6, value="=50+50")
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def xls_file(plan):
    book = xlwt.Workbook()
    sheet = book.add_sheet("plan")
    for column, value in enumerate(HEADERS):
        sheet.write(0, column, value)
    for row_number, row in enumerate(rows_for(plan), start=1):
        for column, value in enumerate(row):
            sheet.write(row_number, column, value)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


@pytest.mark.parametrize(("filename", "content_type", "builder"), [
    ("plan.csv", "text/csv", csv_file),
    ("plan.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", xlsx_file),
    ("plan.xls", "application/vnd.ms-excel", xls_file),
])
def test_tabular_plan_formats_import_the_same_plan(client, plan, filename, content_type, builder):
    response = client.post("/api/plans/import", files={"file": (filename, builder(plan), content_type)})
    assert response.status_code == 200, response.text
    imported = response.json()
    assert imported["plan"] == {
        "plan_id": plan["plan_id"],
        "scenario_id": plan["scenario_id"],
        "decisions": plan["decisions"],
    }
    assert imported["normalized"] is False
    assert imported["model_checked"] is False


def test_semicolon_csv_accepts_decimal_comma(client, plan):
    content = csv_file(plan, delimiter=";").decode("utf-8").replace("12.33", "12,33")
    response = client.post("/api/plans/import", files={"file": ("план.csv", content.encode(), "text/csv")})
    assert response.status_code == 200, response.text
    assert response.json()["plan"]["decisions"]["inventory_policy"]["initial_inventory_t"] == 12.33


def test_comma_delimited_csv_rejects_ambiguous_decimal_comma(client):
    content = ",".join(HEADERS) + '\nplan,p,BASE,,,,,,\ninventory_policy,,,,,,,,"12,33"\n'
    response = client.post("/api/plans/import", files={"file": ("plan.csv", content.encode())})
    assert response.status_code == 422, response.text
    assert "требуется число" in str(response.json()["detail"])


@pytest.mark.parametrize(("format", "content_type"), [
    ("csv", "text/csv; charset=utf-8"),
    ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
])
def test_downloaded_template_is_a_valid_import(client, format, content_type):
    template = client.get("/api/plans/import/template", params={"format": format})
    assert template.status_code == 200, template.text
    assert template.headers["content-type"] == content_type
    imported = client.post("/api/plans/import", files={"file": (f"template.{format}", template.content, content_type)})
    assert imported.status_code == 200, imported.text
    assert imported.json()["plan"]["decisions"]["supply_orders"] == []


@pytest.mark.parametrize(("filename", "contents", "cause"), [
    ("plan.xlsx", b"not a zip", "сигнатура"),
    ("plan.xls", b"not an ole file", "сигнатура"),
    ("plan.csv", b"record_type,unknown\nplan,value\n", "Неизвестные колонки"),
])
def test_tabular_import_rejects_invalid_files(client, filename, contents, cause):
    response = client.post("/api/plans/import", files={"file": (filename, contents)})
    assert response.status_code in (415, 422), response.text
    assert cause.lower() in str(response.json()["detail"]).lower()


def test_xlsx_formulas_are_rejected(client, plan):
    response = client.post("/api/plans/import", files={"file": ("plan.xlsx", xlsx_file(plan, formula=True))})
    assert response.status_code == 422, response.text
    assert "формул" in str(response.json()["detail"]).lower()


def test_tabular_validation_reports_source_row(client):
    content = "record_type,plan_id,scenario_id,year,source_id,order_t,reserved_capacity_t,investment_id,initial_inventory_t\nplan,p,BASE,,,,,,\nsupply_order,,,2035,A,-1,,,\ninventory_policy,,,,,,,,12.33\n"
    response = client.post("/api/plans/import", files={"file": ("plan.csv", content.encode())})
    assert response.status_code == 422, response.text
    assert "строка 3" in str(response.json()["detail"]).lower()


def test_csv_never_ignores_a_value_without_a_header(client):
    content = ",".join(HEADERS) + "\nplan,p,BASE,,,,,,,hidden\n"
    response = client.post("/api/plans/import", files={"file": ("plan.csv", content.encode())})
    assert response.status_code == 422, response.text
    assert "без заголовка" in str(response.json()["detail"])


def test_unknown_record_type_reports_its_source_row(client):
    content = ",".join(HEADERS) + "\nplan,p,BASE,,,,,,\nmystery,,,,,,,,\ninventory_policy,,,,,,,,12.33\n"
    response = client.post("/api/plans/import", files={"file": ("plan.csv", content.encode())})
    assert response.status_code == 422, response.text
    assert "Строка 3" in str(response.json()["detail"])
    assert "mystery" in str(response.json()["detail"])

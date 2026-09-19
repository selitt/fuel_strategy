"""Bounded CSV/XLSX/XLS transport for participant plans."""
from __future__ import annotations

import csv
from io import BytesIO, StringIO
import math
from pathlib import PurePosixPath
from typing import Any, Iterable
from zipfile import BadZipFile, ZipFile

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
import xlrd

from src.app.service.json_io import InvalidJSON, loads


PLAN_COLUMNS = (
    "record_type", "plan_id", "scenario_id", "year", "source_id",
    "order_t", "reserved_capacity_t", "investment_id", "initial_inventory_t",
    "delivery_date", "batch_id",
)
ALLOWED_FORMATS = ("json", "csv", "xlsx", "xls")
MAX_TABLE_ROWS = 10_000
MAX_CELL_CHARS = 10_000
MAX_XLSX_MEMBERS = 128
MAX_XLSX_UNCOMPRESSED_BYTES = 16 * 1024 * 1024
XLS_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")


def _problem(status: int, message: str) -> HTTPException:
    return HTTPException(status_code=status, detail=message)


def _name_and_suffix(filename: str | None) -> tuple[str, str]:
    name = (filename or "").replace("\\", "/")
    suffix = PurePosixPath(name).suffix.lower().removeprefix(".")
    if suffix not in ALLOWED_FORMATS:
        raise _problem(415, "Поддерживаются файлы JSON, CSV, XLSX и XLS")
    return name, suffix


def _check_cell(value: Any, row_number: int, column: str) -> Any:
    if isinstance(value, str):
        value = value.strip()
        if len(value) > MAX_CELL_CHARS:
            raise _problem(422, f"Строка {row_number}, колонка {column}: значение слишком длинное")
        if value.startswith("="):
            raise _problem(422, f"Строка {row_number}, колонка {column}: формулы запрещены")
    return value


def _csv_rows(contents: bytes) -> tuple[list[tuple[int, list[Any]]], bool]:
    try:
        text = contents.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise _problem(422, "CSV должен быть сохранён в UTF-8") from exc
    if "\x00" in text:
        raise _problem(422, "CSV содержит недопустимый нулевой байт")
    sample = text[:8192]
    delimiter = next((candidate for candidate in (",", ";")
                      if "record_type" in next(csv.reader(StringIO(sample), delimiter=candidate), [])), None)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;").delimiter
        except csv.Error as exc:
            raise _problem(422, "Не удалось определить разделитель CSV: используйте запятую или точку с запятой") from exc
    try:
        rows = [(number, row) for number, row in enumerate(csv.reader(StringIO(text), delimiter=delimiter), start=1)]
        return rows, delimiter == ";"
    except csv.Error as exc:
        raise _problem(422, "CSV имеет некорректную структуру") from exc


def _check_xlsx_archive(contents: bytes) -> None:
    if not contents.startswith(b"PK\x03\x04"):
        raise _problem(415, "Сигнатура файла не соответствует формату XLSX")
    try:
        with ZipFile(BytesIO(contents)) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_MEMBERS:
                raise _problem(422, "XLSX содержит слишком много внутренних файлов")
            if any(member.flag_bits & 0x1 for member in members):
                raise _problem(422, "Зашифрованные XLSX не поддерживаются")
            if sum(member.file_size for member in members) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise _problem(413, "Распакованный XLSX превышает разрешённый размер")
    except BadZipFile as exc:
        raise _problem(422, "XLSX повреждён или имеет некорректную ZIP-структуру") from exc


def _xlsx_rows(contents: bytes) -> list[tuple[int, list[Any]]]:
    _check_xlsx_archive(contents)
    try:
        book = load_workbook(BytesIO(contents), read_only=True, data_only=False, keep_links=False)
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise _problem(422, "Не удалось прочитать XLSX") from exc
    try:
        if "plan" not in book.sheetnames:
            raise _problem(422, "В книге XLSX отсутствует лист plan")
        sheet = book["plan"]
        if sheet.max_column > len(PLAN_COLUMNS):
            raise _problem(422, f"Лист plan содержит больше {len(PLAN_COLUMNS)} колонок")
        rows = []
        for number, cells in enumerate(sheet.iter_rows(max_col=len(PLAN_COLUMNS)), start=1):
            values = []
            for column, cell in zip(PLAN_COLUMNS, cells):
                if cell.data_type == "f":
                    raise _problem(422, f"Строка {number}, колонка {column}: формулы запрещены")
                values.append(cell.value)
            rows.append((number, values))
            if len(rows) > MAX_TABLE_ROWS:
                raise _problem(413, f"Таблица содержит больше {MAX_TABLE_ROWS} строк")
        return rows
    finally:
        book.close()


def _xls_rows(contents: bytes) -> list[tuple[int, list[Any]]]:
    if not contents.startswith(XLS_SIGNATURE):
        raise _problem(415, "Сигнатура файла не соответствует формату XLS")
    try:
        book = xlrd.open_workbook(file_contents=contents, on_demand=True)
        sheet = book.sheet_by_name("plan")
    except (xlrd.XLRDError, IndexError) as exc:
        raise _problem(422, "Не удалось прочитать XLS или в книге отсутствует лист plan") from exc
    try:
        if sheet.nrows > MAX_TABLE_ROWS:
            raise _problem(413, f"Таблица содержит больше {MAX_TABLE_ROWS} строк")
        if sheet.ncols > len(PLAN_COLUMNS):
            raise _problem(422, f"Лист plan содержит больше {len(PLAN_COLUMNS)} колонок")
        rows = []
        for index in range(sheet.nrows):
            values = []
            for column in range(len(PLAN_COLUMNS)):
                if column >= sheet.ncols:
                    values.append(None)
                    continue
                cell = sheet.cell(index, column)
                if cell.ctype == xlrd.XL_CELL_ERROR:
                    raise _problem(422, f"Строка {index + 1}, колонка {PLAN_COLUMNS[column]}: ошибка Excel")
                values.append(None if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) else cell.value)
            rows.append((index + 1, values))
        return rows
    finally:
        book.release_resources()


def _number(
    value: Any, row_number: int, column: str, *, integer: bool = False, allow_decimal_comma: bool = False,
) -> int | float:
    if isinstance(value, bool) or value in (None, ""):
        raise _problem(422, f"Строка {row_number}, колонка {column}: требуется число")
    normalized = value.replace(",", ".") if allow_decimal_comma and isinstance(value, str) else value
    try:
        number = float(normalized)
    except (TypeError, ValueError) as exc:
        raise _problem(422, f"Строка {row_number}, колонка {column}: требуется число") from exc
    if not math.isfinite(number):
        raise _problem(422, f"Строка {row_number}, колонка {column}: требуется конечное число")
    if number < 0:
        raise _problem(422, f"Строка {row_number}, колонка {column}: значение не может быть отрицательным")
    if integer:
        if not number.is_integer():
            raise _problem(422, f"Строка {row_number}, колонка {column}: требуется целое число")
        return int(number)
    return number


def _required_text(row: dict[str, Any], row_number: int, column: str) -> str:
    value = row[column]
    if value in (None, ""):
        raise _problem(422, f"Строка {row_number}, колонка {column}: значение обязательно")
    return str(value).strip()


def _reject_unused(row: dict[str, Any], row_number: int, allowed: set[str]) -> None:
    unexpected = [column for column, value in row.items() if column not in allowed and value not in (None, "")]
    if unexpected:
        raise _problem(422, f"Строка {row_number}: поля {', '.join(unexpected)} не применимы к типу {row['record_type']}")


def _table_to_plan(
    source_rows: Iterable[tuple[int, list[Any]]], *, allow_decimal_comma: bool = False,
) -> dict[str, Any]:
    rows = [(number, values) for number, values in source_rows if any(value not in (None, "") for value in values)]
    if not rows:
        raise _problem(422, "Таблица плана пуста")
    header_number, header_values = rows[0]
    headers = [str(value).strip() if value is not None else "" for value in header_values]
    while headers and not headers[-1]:
        headers.pop()
    if len(headers) != len(set(headers)):
        raise _problem(422, f"Строка {header_number}: заголовки колонок не должны повторяться")
    unknown = sorted(set(headers) - set(PLAN_COLUMNS) - {""})
    missing = sorted(set(PLAN_COLUMNS) - {"delivery_date", "batch_id"} - set(headers))
    if unknown:
        raise _problem(422, f"Неизвестные колонки: {', '.join(unknown)}")
    if missing:
        raise _problem(422, f"Отсутствуют обязательные колонки: {', '.join(missing)}")
    if len(rows) - 1 > MAX_TABLE_ROWS:
        raise _problem(413, f"Таблица содержит больше {MAX_TABLE_ROWS} строк")

    plan_rows: list[tuple[int, dict[str, Any]]] = []
    for number, values in rows[1:]:
        if len(values) > len(headers) and any(value not in (None, "") for value in values[len(headers):]):
            raise _problem(422, f"Строка {number}: найдено значение без заголовка колонки")
        padded = [*values, *([None] * (len(headers) - len(values)))]
        row = {header: _check_cell(padded[index], number, header) for index, header in enumerate(headers) if header}
        row_type = str(row.get("record_type") or "").strip()
        if not row_type:
            raise _problem(422, f"Строка {number}, колонка record_type: значение обязательно")
        row["record_type"] = row_type
        plan_rows.append((number, row))

    grouped: dict[str, list[tuple[int, dict[str, Any]]]] = {}
    for number, row in plan_rows:
        grouped.setdefault(row["record_type"], []).append((number, row))
    unknown_types = sorted(set(grouped) - {"plan", "inventory_policy", "supply_order", "capacity_reservation", "investment"})
    if unknown_types:
        number, row = next((number, row) for number, row in plan_rows if row["record_type"] in unknown_types)
        raise _problem(422, f"Строка {number}: неизвестный record_type {row['record_type']}")
    for required in ("plan", "inventory_policy"):
        count = len(grouped.get(required, []))
        if count != 1:
            raise _problem(422, f"Требуется ровно одна строка record_type={required}; найдено {count}")

    number, metadata = grouped["plan"][0]
    _reject_unused(metadata, number, {"record_type", "plan_id", "scenario_id"})
    plan = {
        "plan_id": _required_text(metadata, number, "plan_id"),
        "scenario_id": _required_text(metadata, number, "scenario_id"),
        "decisions": {"supply_orders": [], "capacity_reservations": [], "investments": [], "inventory_policy": {}},
    }
    number, inventory = grouped["inventory_policy"][0]
    _reject_unused(inventory, number, {"record_type", "initial_inventory_t"})
    plan["decisions"]["inventory_policy"] = {
        "initial_inventory_t": _number(
            inventory["initial_inventory_t"], number, "initial_inventory_t", allow_decimal_comma=allow_decimal_comma,
        ),
    }
    for number, row in grouped.get("supply_order", []):
        _reject_unused(row, number, {"record_type", "year", "source_id", "order_t", "delivery_date", "batch_id"})
        plan["decisions"]["supply_orders"].append({
            "year": _number(row["year"], number, "year", integer=True, allow_decimal_comma=allow_decimal_comma),
            "source_id": _required_text(row, number, "source_id"),
            "order_t": _number(row["order_t"], number, "order_t", allow_decimal_comma=allow_decimal_comma),
            **{key: str(row[key]) for key in ("delivery_date", "batch_id") if row.get(key) not in (None, "")},
        })
    for number, row in grouped.get("capacity_reservation", []):
        _reject_unused(row, number, {"record_type", "year", "source_id", "reserved_capacity_t"})
        plan["decisions"]["capacity_reservations"].append({
            "year": _number(row["year"], number, "year", integer=True, allow_decimal_comma=allow_decimal_comma),
            "source_id": _required_text(row, number, "source_id"),
            "reserved_capacity_t": _number(
                row["reserved_capacity_t"], number, "reserved_capacity_t", allow_decimal_comma=allow_decimal_comma,
            ),
        })
    for number, row in grouped.get("investment", []):
        _reject_unused(row, number, {"record_type", "year", "investment_id"})
        plan["decisions"]["investments"].append({
            "investment_id": _required_text(row, number, "investment_id"),
            "year": _number(row["year"], number, "year", integer=True, allow_decimal_comma=allow_decimal_comma),
        })
    return plan


def parse_plan_file(contents: bytes, filename: str | None, limit: int) -> Any:
    """Parse a bounded plan upload based on its declared extension and signature."""
    _, suffix = _name_and_suffix(filename)
    if len(contents) > limit:
        raise _problem(413, "Файл превышает разрешённый размер")
    if suffix == "json":
        try:
            return loads(contents)
        except InvalidJSON as exc:
            raise _problem(422, str(exc)) from exc
    if suffix == "csv":
        rows, allow_decimal_comma = _csv_rows(contents)
        return _table_to_plan(rows, allow_decimal_comma=allow_decimal_comma)
    if suffix == "xlsx":
        return _table_to_plan(_xlsx_rows(contents))
    return _table_to_plan(_xls_rows(contents))


def _template_rows() -> list[list[Any]]:
    return [
        list(PLAN_COLUMNS),
        ["plan", "plan-import", "BASE", "", "", "", "", "", ""],
        ["inventory_policy", "", "", "", "", "", "", "", 0],
    ]


def plan_template(format: str) -> tuple[bytes, str, str]:
    if format == "csv":
        output = StringIO()
        csv.writer(output).writerows(_template_rows())
        return ("\ufeff" + output.getvalue()).encode("utf-8"), "text/csv; charset=utf-8", "plan-template.csv"
    if format == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = "plan"
        for row in _template_rows():
            sheet.append(row)
        output = BytesIO()
        book.save(output)
        return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "plan-template.xlsx"
    raise _problem(422, "Шаблон доступен только в форматах csv и xlsx")
